"""Turn files into an ordered list of text segments (headings + body) with page numbers.

One parser is shared by folder search and single-document chat, so both see
exactly the same text.
"""
from __future__ import annotations

import logging
import re
import threading
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Callable, List, Optional

if TYPE_CHECKING:
    from .ocr import OcrEngine

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {"pdf", "docx", "txt", "md"}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "tif", "tiff", "bmp", "webp"}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | IMAGE_EXTENSIONS

MIN_PAGE_CHARS = 25     # a PDF page with fewer non-space characters has no usable text layer
# Pages whose content is mostly a picture (scans, screenshots of notes) often still carry a
# little real text: a watermark, a question number. Measured on a 66-page question bank:
# ~25 such pages had 34-190 characters of text but images covering 37-71% of the page.
IMAGE_PAGE_COVERAGE = 0.3   # images covering this much of a page usually hold content
LOW_TEXT_CHARS = 250        # ...and with this little text the page is effectively a scan
REGION_MIN_COVERAGE = 0.08  # on real text pages, OCR images at least this big (diagrams, screenshots)
OCR_DPI = 150           # measured on the sample paper: 99.8% word recall, same as 300 DPI
MAX_IMAGE_SIDE = 3000   # larger photos are downscaled before OCR (it caps at 2000 px anyway)
MAX_FRAMES = 500        # multi-page TIFF safety limit

# pdfium is not thread-safe, and the background indexer and a "chat with this
# file" request can parse PDFs at the same time. Every pdfium call holds this.
_PDFIUM_LOCK = threading.RLock()

ProgressFn = Callable[[str], None]


class ParseError(Exception):
    """The file could not be read (corrupt, encrypted, unsupported)."""


@dataclass
class Segment:
    text: str
    page: Optional[int] = None  # 1-based; None for formats without pages
    is_heading: bool = False
    ocr: bool = False           # text came from OCR (may contain recognition errors)


@dataclass
class ParsedDocument:
    path: str
    segments: List[Segment] = field(default_factory=list)
    num_pages: Optional[int] = None
    ocr_pages: int = 0    # pages (or images) whose text was recognised with OCR
    ocr_needed: int = 0   # pages that needed OCR but it was unavailable

    @property
    def char_count(self) -> int:
        return sum(len(s.text) for s in self.segments)


# ---------------------------------------------------------------------------
# Heading detection for plain-text lines (PDF / TXT have no style information)
# ---------------------------------------------------------------------------

_NUMBER_PREFIX = re.compile(r"^(?:\d{1,2}(?:\.\d{1,2})*\.?|[IVX]{1,5}\.)\s+")
_KNOWN_SECTIONS = {
    "abstract", "introduction", "background", "related work", "prior work",
    "method", "methods", "methodology", "approach", "model", "architecture",
    "experiments", "experiment", "experimental setup", "experimental results",
    "evaluation", "results", "results and discussion", "discussion", "analysis",
    "conclusion", "conclusions", "future work", "limitations", "summary",
    "references", "bibliography", "acknowledgements", "acknowledgments", "appendix",
}


def is_heading_line(line: str) -> bool:
    s = line.strip()
    if not 3 <= len(s) <= 90:
        return False
    letters = [c for c in s if c.isalpha()]
    if len(letters) < 3 or len(s.split()) > 12:
        return False
    if s[-1] in ".,;":
        return False
    # Number-heavy lines are table rows, footers or equations, not headings.
    letter_ratio = len(letters) / len([c for c in s if not c.isspace()])
    if letter_ratio < 0.5:
        return False
    core = _NUMBER_PREFIX.sub("", s)
    if core.lower().rstrip(":") in _KNOWN_SECTIONS:
        return True
    # ALL-CAPS lines ("PROPOSED METHOD"). Case-sensitive on purpose: the old regex
    # used IGNORECASE and matched almost every line of body text. A lone caps word
    # is usually a figure label or acronym ("LSTM"), so require two words.
    # Caps lines with numbers ("ICASSP 2018", "PESQ 1.54") are footers and table cells.
    if (len(letters) >= 4 and len(core.split()) >= 2 and letter_ratio >= 0.7
            and sum(c.isupper() for c in letters) / len(letters) > 0.9):
        return True
    # Numbered title-case headings ("3.2 Network Architecture").
    if _NUMBER_PREFIX.match(s) and core[:1].isupper() and len(core.split()) <= 8:
        return True
    return False


# ---------------------------------------------------------------------------
# Text clean-up helpers
# ---------------------------------------------------------------------------

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f￾￿]")


def _clean(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    # pdfium marks a soft (line-break) hyphen as \x02: "recon\x02\nstructing".
    text = re.sub(r"\x02\s*\n\s*", "", text)
    text = _CONTROL.sub("", text)
    return text


def _join_lines(lines: List[str], ocr: bool = False) -> str:
    """Join wrapped lines of one paragraph back into prose."""
    out = ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if not out:
            out = line
        elif ocr and out.endswith("-") and line[:1].islower():
            out = out[:-1] + line  # scans have no soft-hyphen marker: "recon-" + "struction"
        elif out.endswith("-") and line[:1].isalpha():
            out += line  # "state-of-the-" + "art"
        else:
            out += " " + line
    return out


def _segments_from_lines(lines: List[str], page: Optional[int], ocr: bool = False) -> List[Segment]:
    segments: List[Segment] = []
    body: List[str] = []

    def flush():
        text = _join_lines(body, ocr)
        if text:
            segments.append(Segment(text=text, page=page, ocr=ocr))
        body.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()  # block break (e.g. the end of an OCR column)
            continue
        if is_heading_line(stripped):
            flush()
            segments.append(Segment(text=stripped, page=page, is_heading=True, ocr=ocr))
        else:
            body.append(stripped)
    flush()
    return segments


def _ocr_lines_safely(ocr: "OcrEngine"):
    def run(image) -> List[str]:
        try:
            return [_clean(line) for line in ocr.read_lines(image)]
        except Exception as e:  # one bad page must not sink the whole document
            logger.warning("OCR failed on a page: %s", e)
            return []
    return run


# ---------------------------------------------------------------------------
# Format readers
# ---------------------------------------------------------------------------

def _repeated_margin_lines(pages: List[List[str]]) -> set:
    """Running headers/footers: lines that appear at a page edge on most pages."""
    if len(pages) < 3:
        return set()
    counts: Counter = Counter()
    for lines in pages:
        edge = [l.strip() for l in lines[:2] + lines[-2:] if l.strip()]
        counts.update(set(edge))
    threshold = max(3, len(pages) // 2)
    return {line for line, n in counts.items() if n >= threshold}


def _image_boxes(page, pdfium_c) -> List[tuple]:
    """Bounding boxes (PDF points) of image objects on a page, clipped to the page."""
    width, height = page.get_size()
    boxes = []
    for obj in page.get_objects(filter=[pdfium_c.FPDF_PAGEOBJ_IMAGE], max_depth=3):
        left, bottom, right, top = obj.get_bounds()
        left, right = max(left, 0.0), min(right, width)
        bottom, top = max(bottom, 0.0), min(top, height)
        if right > left and top > bottom:
            boxes.append((left, bottom, right, top))
    return boxes


def _ocr_plan(chars: int, boxes: List[tuple], width: float, height: float, rotated: bool):
    """What to OCR on a page: ("full", None), ("regions", [boxes]) or (None, None)."""
    if not boxes:
        return None, None
    page_area = width * height
    coverage = min(1.0, sum((r - l) * (t - b) for l, b, r, t in boxes) / page_area)
    if chars < MIN_PAGE_CHARS or (chars < LOW_TEXT_CHARS and coverage >= IMAGE_PAGE_COVERAGE):
        return "full", None  # a scan, maybe with a watermark or a question number on top
    if coverage >= IMAGE_PAGE_COVERAGE and not rotated:
        big = [box for box in boxes if (box[2] - box[0]) * (box[3] - box[1]) / page_area >= REGION_MIN_COVERAGE]
        if big:
            return "regions", big  # real text plus large pictures: keep the text, read the pictures
    return None, None


def _read_pdf(path: str, ocr: Optional["OcrEngine"], progress: Optional[ProgressFn]) -> ParsedDocument:
    import pypdfium2 as pdfium
    import pypdfium2.raw as pdfium_c

    with _PDFIUM_LOCK:
        try:
            pdf = pdfium.PdfDocument(path)
        except pdfium.PdfiumError as e:
            raise ParseError(f"cannot open PDF ({e})") from e

    scale = OCR_DPI / 72
    try:
        page_lines: List[List[str]] = []
        plans = {}  # page index -> (kind, boxes, page height)
        with _PDFIUM_LOCK:
            for i in range(len(pdf)):
                page = pdf[i]
                textpage = page.get_textpage()
                try:
                    lines = _clean(textpage.get_text_bounded()).split("\n")
                    page_lines.append(lines)
                    width, height = page.get_size()
                    kind, boxes = _ocr_plan(len(re.sub(r"\s", "", "".join(lines))), _image_boxes(page, pdfium_c),
                                            width, height, rotated=page.get_rotation() != 0)
                    if kind:
                        plans[i] = (kind, boxes, height)
                finally:
                    textpage.close()
                    page.close()

        full_ocr, region_ocr = {}, {}
        if plans and ocr is not None and ocr.available:
            def render(i: int):
                kind, boxes, height = plans[i]
                with _PDFIUM_LOCK:  # render under the lock; OCR itself runs unlocked, in parallel
                    page = pdf[i]
                    bitmap = page.render(scale=scale)  # BGR, which is what RapidOCR expects
                    image = bitmap.to_numpy().copy()
                    bitmap.close()
                    page.close()
                if kind == "full":
                    return [image]
                # PDF y runs bottom-up; pixel rows run top-down.
                return [image[int((height - t) * scale):int((height - b) * scale) + 1,
                              int(l * scale):int(r * scale) + 1] for l, b, r, t in boxes]

            read = _ocr_lines_safely(ocr)

            def run(i: int) -> List[str]:
                lines: List[str] = []
                for crop in render(i):
                    if crop.size and min(crop.shape[:2]) >= 16:
                        lines += read(crop) + [""]
                return lines

            pages = sorted(plans)
            for n, (i, lines) in enumerate(zip(pages, ocr.map_ordered(run, pages)), start=1):
                (full_ocr if plans[i][0] == "full" else region_ocr)[i] = lines
                if progress:
                    progress(f"OCR page {n} of {len(pages)}")
    finally:
        with _PDFIUM_LOCK:
            pdf.close()

    for i, lines in full_ocr.items():
        page_lines[i] = lines
    margins = _repeated_margin_lines(page_lines)
    doc = ParsedDocument(path=path, num_pages=len(page_lines), ocr_pages=len(full_ocr) + len(region_ocr),
                         ocr_needed=len(plans) - len(full_ocr) - len(region_ocr))

    def keep(lines: List[str]) -> List[str]:
        return [l for l in lines if not l.strip() or (l.strip() not in margins and not l.strip().isdigit())]

    for i, lines in enumerate(page_lines):
        doc.segments.extend(_segments_from_lines(keep(lines), page=i + 1, ocr=i in full_ocr))
        if i in region_ocr:  # text read from pictures on an otherwise digital page
            doc.segments.extend(_segments_from_lines(keep(region_ocr[i]), page=i + 1, ocr=True))
    return doc


def _read_image(path: str, ocr: Optional["OcrEngine"], progress: Optional[ProgressFn]) -> ParsedDocument:
    from PIL import Image, ImageOps, ImageSequence

    try:
        image = Image.open(path)
        frame_count = getattr(image, "n_frames", 1)
    except Exception as e:
        raise ParseError(f"cannot open image ({e})") from e

    multi = frame_count > 1
    doc = ParsedDocument(path=path, num_pages=frame_count if multi else None)
    if ocr is None or not ocr.available:
        image.close()
        doc.ocr_needed = frame_count
        return doc

    def frames():
        for k, frame in enumerate(ImageSequence.Iterator(image)):
            if k >= MAX_FRAMES:
                break
            frame = ImageOps.exif_transpose(frame).convert("RGB")  # phone photos come rotated
            if max(frame.size) > MAX_IMAGE_SIDE:
                frame.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
            yield frame

    try:
        results = ocr.map_ordered(_ocr_lines_safely(ocr), frames())
        for k, lines in enumerate(results):
            doc.segments.extend(_segments_from_lines(lines, page=k + 1 if multi else None, ocr=True))
            doc.ocr_pages += 1
            if progress and multi:
                progress(f"OCR page {k + 1} of {frame_count}")
    except OSError as e:  # truncated or corrupt image data
        raise ParseError(f"cannot read image ({e})") from e
    finally:
        image.close()
    return doc


def _read_docx(path: str) -> ParsedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(path)
    except Exception as e:  # python-docx raises several unrelated types
        raise ParseError(f"cannot open DOCX ({e})") from e

    doc = ParsedDocument(path=path)
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = _clean(para.text).strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            heading = style.startswith("Heading") or style == "Title"
            doc.segments.append(Segment(text=text, is_heading=heading))
        elif tag == "tbl":
            table = Table(child, document)
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    value = _clean(cell.text).strip().replace("\n", " ")
                    if value and (not cells or cells[-1] != value):  # merged cells repeat
                        cells.append(value)
                if cells:
                    rows.append(" | ".join(cells))
            if rows:
                doc.segments.append(Segment(text="\n".join(rows)))
    return doc


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


_MD_HEADING = re.compile(r"^\s{0,3}#{1,6}\s+(.*?)\s*#*\s*$")


def _read_text(path: str, markdown: bool) -> ParsedDocument:
    with open(path, "rb") as fh:
        text = _clean(_decode(fh.read()))

    doc = ParsedDocument(path=path)
    for block in re.split(r"\n\s*\n", text):
        lines = [l for l in block.split("\n") if l.strip()]
        if not lines:
            continue
        if markdown:
            body: List[str] = []
            for line in lines:
                m = _MD_HEADING.match(line)
                if m:
                    if body:
                        doc.segments.append(Segment(text=_join_lines(body)))
                        body = []
                    doc.segments.append(Segment(text=m.group(1), is_heading=True))
                else:
                    body.append(line)
            if body:
                doc.segments.append(Segment(text=_join_lines(body)))
        elif len(lines) == 1 and is_heading_line(lines[0]):
            doc.segments.append(Segment(text=lines[0].strip(), is_heading=True))
        else:
            doc.segments.append(Segment(text=_join_lines(lines)))
    return doc


def parse_file(path: str, ocr: Optional["OcrEngine"] = None,
               progress: Optional[ProgressFn] = None) -> ParsedDocument:
    """Parse a file. With `ocr`, scanned PDF pages and image files are run through OCR."""
    ext = Path(path).suffix.lstrip(".").lower()
    if ext == "pdf":
        return _read_pdf(path, ocr, progress)
    if ext in IMAGE_EXTENSIONS:
        return _read_image(path, ocr, progress)
    if ext == "docx":
        return _read_docx(path)
    if ext in ("txt", "md"):
        return _read_text(path, markdown=(ext == "md"))
    raise ParseError(f"unsupported file type: .{ext}")
