"""Questions about positions in a document ("the last page", "page 3") and image-heavy pages."""
import shutil

import pytest

from doclamar.engine import Engine
from doclamar.parsing import _ocr_plan
from doclamar.pipeline import page_reference
from doclamar.store import Scope
from tests.conftest import SAMPLE_PDF, FakeEmbedder, FakeLLM, FakeOcr, FakeReranker


@pytest.mark.parametrize("q, expected", [
    ("what does the last page contain", [-1]), ("whats on page 22", [22]), ("summarize pages 3 to 5", [3, 4, 5]),
    ("explain pages 3-4", [3, 4]), ("what is on the first page", [1]), ("the 5th page", [5]),
    ("second to last page", [-2]), ("last 3 pages", [-3, -2, -1]), ("first two pages", [1, 2]),
    ("page no. 7", [7]), ("pages 4 and 9", [4, 9]), ("title page", [1]),
    ("explain page tables in 80386", None), ("what is paging", None), ("how many pages does it have", None),
    ("last stage of the pipeline", None),
])
def test_page_reference(q, expected):
    assert page_reference(q) == expected


def test_ocr_plan():
    page = (612.0, 792.0)
    big = (50, 50, 560, 500)        # ~47% of the page
    small = (50, 700, 120, 760)     # a logo
    assert _ocr_plan(34, [big], *page, rotated=False) == ("full", None)       # scan + watermark text
    assert _ocr_plan(0, [small], *page, rotated=False) == ("full", None)      # no text at all
    assert _ocr_plan(600, [big, small], *page, rotated=False) == ("regions", [big])  # text + diagram
    assert _ocr_plan(900, [small], *page, rotated=False) == (None, None)      # text + logo
    assert _ocr_plan(900, [], *page, rotated=False) == (None, None)


def _engine(settings, llm):
    return Engine(settings, embedder=FakeEmbedder(), reranker=FakeReranker(), llm=llm, ocr=FakeOcr())


@pytest.fixture
def setup(settings, tmp_path):
    import pypdfium2 as pdfium
    from docx import Document

    folder = tmp_path / "docs"
    folder.mkdir()
    shutil.copy(SAMPLE_PDF, folder / "paper.pdf")  # 5 pages
    # 3 pages: two from the paper, then a blank page (no text, no images -> nothing to read)
    src = pdfium.PdfDocument(str(SAMPLE_PDF))
    mixed = pdfium.PdfDocument.new()
    mixed.import_pages(src, [0, 1])
    mixed.new_page(612, 792)
    mixed.save(str(folder / "with_blank.pdf"))
    mixed.close()
    src.close()
    d = Document()
    d.add_paragraph("Opening paragraph of the memo.")
    d.add_paragraph("Closing paragraph: the budget was approved.")
    d.save(folder / "memo.docx")

    llm = FakeLLM(answer="It lists the references [1].")
    e = _engine(settings, llm)
    e.indexer.sync_folder(str(folder))
    yield e, llm, folder
    e.close()


def test_last_page_question_reads_the_last_page(setup):
    e, llm, folder = setup
    result = e.answer("what does the last page contain", Scope(files=[str(folder / "paper.pdf")]))
    assert result["status"] == "answered"
    assert result["citations"] and all(c["page"] in (4, 5) for c in result["citations"])
    assert any(c["page"] == 5 or c["pages"].endswith("5") for c in result["citations"])
    assert "page 5 of paper.pdf, which has 5 pages" in llm.prompts[-1]
    assert "Document: paper.pdf, 5 pages." in llm.prompts[-1]


def test_page_range_and_ordinal(setup):
    e, llm, folder = setup
    scope = Scope(files=[str(folder / "paper.pdf")])
    pages = {c["page"] for c in e.answer("summarize pages 2 to 3", scope)["citations"]}
    assert pages and pages <= {1, 2, 3}  # a chunk may start on the page before
    assert {c["page"] for c in e.answer("what is on the first page", scope)["citations"]} == {1}


def test_missing_page_is_answered_without_the_llm(setup):
    e, llm, folder = setup
    calls = len(llm.prompts)
    result = e.answer("what is on page 9", Scope(files=[str(folder / "paper.pdf")]))
    assert result["answer"] == "paper.pdf has 5 pages, so there is no page 9."
    assert len(llm.prompts) == calls


def test_blank_page_says_so(setup):
    e, llm, folder = setup
    result = e.answer("what does the last page contain", Scope(files=[str(folder / "with_blank.pdf")]))
    assert "Page 3 of with_blank.pdf has no readable text" in result["answer"]


def test_files_without_pages(setup):
    e, llm, folder = setup
    scope = Scope(files=[str(folder / "memo.docx")])
    last = e.answer("what is on the last page", scope)
    assert last["citations"] and "Closing paragraph" in last["citations"][-1]["snippet"]
    assert "has no page numbers" in e.answer("what is on page 3", scope)["answer"]


def test_page_question_in_a_folder_picks_the_matching_file(setup):
    e, llm, folder = setup
    result = e.answer("what is on page 2 of the lip reading speech reconstruction paper", Scope(folder=str(folder)))
    assert {c["file"] for c in result["citations"]} == {"paper.pdf"}
    assert {c["page"] for c in result["citations"]} <= {1, 2}
