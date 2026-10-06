from __future__ import annotations

import hashlib
import re
import sys
from pathlib import Path
from typing import List, Optional

import numpy as np
import pytest

BACKEND = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(BACKEND))

from doclamar.config import Settings  # noqa: E402
from doclamar.engine import Engine  # noqa: E402

SAMPLE_PDF = BACKEND / "eval" / "data" / "machinelearning.pdf"


def _words(text: str) -> List[str]:
    return re.findall(r"[a-z0-9]+", text.lower())


class FakeEmbedder:
    """Hashed bag-of-words vectors: deterministic, offline, good enough to rank by word overlap."""

    dim = 384

    def __init__(self):
        self.calls = 0

    def _vec(self, text: str) -> np.ndarray:
        v = np.zeros(self.dim, dtype=np.float32)
        for w in _words(text):
            v[int(hashlib.md5(w.encode()).hexdigest(), 16) % self.dim] += 1.0
        n = np.linalg.norm(v)
        return v / n if n else v

    def embed_documents(self, texts, batch_size: int = 32) -> np.ndarray:
        self.calls += 1
        if not texts:
            return np.zeros((0, self.dim), dtype=np.float32)
        return np.vstack([self._vec(t) for t in texts])

    def embed_query(self, query: str) -> np.ndarray:
        return self._vec(query)


class FakeReranker:
    """Logit = 4 * (shared words) - 8, so unrelated passages land near zero relevance."""

    def __init__(self, fixed: Optional[float] = None):
        self.fixed = fixed

    def score(self, query: str, passages) -> List[float]:
        if self.fixed is not None:
            return [self.fixed] * len(passages)
        q = set(_words(query))
        return [4.0 * len(q & set(_words(p))) - 8.0 for p in passages]


class FakeOcr:
    """Stands in for RapidOCR: every page/image 'reads' as the same lines."""

    workers = 1

    def __init__(self, lines=None, available=True, fail=False):
        self.lines = lines if lines is not None else ["SCANNED INVOICE", "Invoice number 4471 for consulting services."]
        self._available = available
        self.fail = fail
        self.calls = 0

    @property
    def available(self) -> bool:
        return self._available

    @property
    def unavailable_reason(self):
        return None if self._available else "OCR engine failed to load: test"

    def read_lines(self, image):
        self.calls += 1
        if self.fail:
            raise RuntimeError("OCR crashed on this page")
        return list(self.lines)

    def map_ordered(self, fn, items):
        return map(fn, items)


def make_scanned_pdf(path: Path, pages=(0,), dpi: int = 150) -> Path:
    """An image-only copy of pages of the sample paper, like a scanner would produce."""
    from eval.scan import make_scanned_pdf as scan

    return scan(SAMPLE_PDF, path, pages=pages, dpi=dpi)


class FakeLLM:
    def __init__(self, configured: bool = True, answer: str = "The answer is in the documents [1]."):
        self._configured = configured
        self.answer = answer
        self.prompts: List[str] = []

    @property
    def configured(self) -> bool:
        return self._configured

    def describe(self) -> dict:
        return {"configured": self._configured, "provider": "fake" if self._configured else None, "model": None}

    def configure(self, provider, api_key, model=None, validate=True):
        self._configured = True
        return self.describe()

    def generate(self, prompt: str, system: str = "", max_tokens: int = 1500, temperature: float = 0.2) -> str:
        self.prompts.append(prompt)
        if prompt.startswith("Rewrite the user's latest message"):
            return "why was the CorrMSE loss function chosen"
        if prompt.startswith("A search of the user's documents"):
            return "rewritten keyword query"
        return self.answer


@pytest.fixture
def settings(tmp_path) -> Settings:
    s = Settings(home=tmp_path / "home")
    s.ensure_dirs()
    return s


@pytest.fixture
def fake_llm() -> FakeLLM:
    return FakeLLM()


@pytest.fixture
def engine(settings, fake_llm):
    e = Engine(settings, embedder=FakeEmbedder(), reranker=FakeReranker(), llm=fake_llm, ocr=FakeOcr())
    yield e
    e.close()


@pytest.fixture
def docs(tmp_path) -> Path:
    """A small folder with each supported format plus things the indexer must skip."""
    root = tmp_path / "docs"
    (root / "sub").mkdir(parents=True)
    (root / "notes.txt").write_text(
        "TRAINING DETAILS\n\nThe model was trained with the CorrMSE loss function. "
        "Optimization used Adam with a learning rate of 0.0001.\n", encoding="utf-8")
    (root / "sub" / "recipe.md").write_text(
        "# Pancakes\n\nWhisk flour, eggs and milk. Cook on a hot pan until golden.\n", encoding="utf-8")
    from docx import Document

    d = Document()
    d.add_heading("Meeting notes", level=1)
    d.add_paragraph("We agreed to evaluate on the GRID corpus next week.")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Owner"
    table.rows[0].cells[1].text = "Shlok"
    d.save(root / "sub" / "meeting.docx")
    # Must be ignored:
    (root / "data.xlsx").write_bytes(b"PK\x03\x04")
    (root / ".hidden").mkdir()
    (root / ".hidden" / "secret.txt").write_text("hidden content", encoding="utf-8")
    (root / "node_modules").mkdir()
    (root / "node_modules" / "readme.txt").write_text("dependency docs", encoding="utf-8")
    (root / "~$lockfile.docx").write_bytes(b"lock")
    return root
